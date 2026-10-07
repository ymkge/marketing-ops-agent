"""PowerPoint deck builder using python-pptx with structured layout."""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from .models import ProposalDeck, SlideContent, ContentBlock


# カラーパレット定義
COLOR_PRIMARY_DARK = RGBColor(15, 23, 42)      # #0F172A (濃紺・タイトル)
COLOR_TEXT_BODY = RGBColor(51, 65, 85)         # #334155 (本文・箇条書き)
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)     # #64748B (補足・注釈)
COLOR_ACCENT = RGBColor(37, 99, 235)           # #2563EB (アクセントブルー)
COLOR_ACCENT_BG = RGBColor(239, 246, 255)      # #EFF6FF (淡いブルー背景)
COLOR_CARD_BG = RGBColor(248, 250, 252)        # #F8FAFC (カード背景)
COLOR_CARD_BORDER = RGBColor(226, 232, 240)    # #E2E8F0 (カード枠線)
COLOR_WHITE = RGBColor(255, 255, 255)


class ProposalPptxBuilder:
    """ProposalDeckモデルを受け取りPowerPointファイルを生成するビルダー"""

    def __init__(self, deck: ProposalDeck):
        self.deck = deck
        self.prs = Presentation()
        # 16:9 ワイド画面設定
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.blank_layout = self.prs.slide_layouts[6]  # 白紙レイアウト

    def build(self, output_path: str) -> str:
        """スライドを生成して指定パスに保存"""
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        
        for slide_data in self.deck.slides:
            if slide_data.slide_type == "cover":
                self._build_cover_slide(slide_data)
            else:
                self._build_content_slide(slide_data)
                
        self.prs.save(output_path)
        return output_path

    def _build_cover_slide(self, slide_data: SlideContent):
        """表紙スライドの生成"""
        slide = self.prs.slides.add_slide(self.blank_layout)

        # 背景上部のアクセント帯
        header_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0), Inches(13.333), Inches(0.2)
        )
        header_bar.fill.solid()
        header_bar.fill.fore_color.rgb = COLOR_ACCENT
        header_bar.line.fill.background()

        # 左側のアクセントバー（縦棒）
        left_accent = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(1.0), Inches(2.2), Inches(0.12), Inches(2.8)
        )
        left_accent.fill.solid()
        left_accent.fill.fore_color.rgb = COLOR_ACCENT
        left_accent.line.fill.background()

        # 提案書タイトル（長さに応じてフォントサイズを動的調整）
        title_box = slide.shapes.add_textbox(
            Inches(1.3), Inches(2.1), Inches(11.0), Inches(1.3)
        )
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = Inches(0)
        tf_title.margin_top = Inches(0)
        p_title = tf_title.paragraphs[0]
        p_title.text = slide_data.title
        # 文字数が多い場合はフォントサイズを調整してサブタイトルとの重なりを防止
        title_len = len(slide_data.title)
        p_title.font.size = Pt(24) if title_len > 45 else (Pt(28) if title_len > 30 else Pt(32))
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_PRIMARY_DARK

        # サブタイトル / リード文
        if slide_data.lead_sentence:
            sub_box = slide.shapes.add_textbox(
                Inches(1.3), Inches(3.5), Inches(11.0), Inches(0.8)
            )
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            tf_sub.margin_left = Inches(0)
            tf_sub.margin_top = Inches(0)
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = slide_data.lead_sentence
            p_sub.font.size = Pt(17)
            p_sub.font.color.rgb = COLOR_TEXT_BODY

        # メタ情報カード（クライアント名、提案者、日付）
        meta_box = slide.shapes.add_textbox(
            Inches(1.3), Inches(5.1), Inches(10.5), Inches(1.5)
        )
        tf_meta = meta_box.text_frame
        tf_meta.word_wrap = True
        tf_meta.margin_left = Inches(0)
        tf_meta.margin_top = Inches(0)

        # クライアント
        p1 = tf_meta.paragraphs[0]
        p1.text = f"宛先: {self.deck.client_name} 御中"
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY_DARK
        p1.space_after = Pt(6)

        # 提案者
        p2 = tf_meta.add_paragraph()
        p2.text = f"提案者: {self.deck.presenter_info or 'マーケティングOps推進チーム'}"
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_TEXT_BODY
        p2.space_after = Pt(6)

        # 日付
        p3 = tf_meta.add_paragraph()
        p3.text = f"提出日: {self.deck.presentation_date or '2026年'}"
        p3.font.size = Pt(12)
        p3.font.color.rgb = COLOR_TEXT_MUTED

        # フッター注記（機密表示）
        if slide_data.footer_note:
            ft_box = slide.shapes.add_textbox(
                Inches(1.3), Inches(6.8), Inches(10.5), Inches(0.4)
            )
            tf_ft = ft_box.text_frame
            p_ft = tf_ft.paragraphs[0]
            p_ft.text = slide_data.footer_note
            p_ft.font.size = Pt(10)
            p_ft.font.color.rgb = COLOR_TEXT_MUTED

    def _build_content_slide(self, slide_data: SlideContent):
        """通常コンテンツスライドの生成（重なり防止レイアウト）"""
        slide = self.prs.slides.add_slide(self.blank_layout)

        # スライド種別のバッジラベル
        badge_labels = {
            "challenges": "01. 課題認識 / SITUATION",
            "target_strategy": "02. ターゲット・訴求設計 / STRATEGY",
            "simulation": "03. 配信シミュレーション / SIMULATION",
            "schedule_team": "04. スケジュール・推進体制 / ROADMAP"
        }
        badge_text = badge_labels.get(slide_data.slide_type, "MARKETING PROPOSAL")

        # バッジ表示
        badge_box = slide.shapes.add_textbox(
            Inches(0.8), Inches(0.45), Inches(6.0), Inches(0.3)
        )
        tf_badge = badge_box.text_frame
        tf_badge.margin_left = Inches(0)
        tf_badge.margin_top = Inches(0)
        p_badge = tf_badge.paragraphs[0]
        p_badge.text = badge_text
        p_badge.font.size = Pt(10.5)
        p_badge.font.bold = True
        p_badge.font.color.rgb = COLOR_ACCENT

        # スライドタイトル
        title_box = slide.shapes.add_textbox(
            Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.6)
        )
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = Inches(0)
        tf_title.margin_top = Inches(0)
        p_title = tf_title.paragraphs[0]
        p_title.text = slide_data.title
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_PRIMARY_DARK

        # リード文（要約・キーメッセージ）のカード背景
        lead_bg = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8), Inches(1.4), Inches(11.73), Inches(0.7)
        )
        lead_bg.fill.solid()
        lead_bg.fill.fore_color.rgb = COLOR_ACCENT_BG
        lead_bg.line.color.rgb = COLOR_CARD_BORDER
        lead_bg.line.width = Pt(0.75)

        tf_lead = lead_bg.text_frame
        tf_lead.word_wrap = True
        tf_lead.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf_lead.margin_left = Inches(0.2)
        tf_lead.margin_right = Inches(0.2)
        p_lead = tf_lead.paragraphs[0]
        p_lead.text = f"POINT: {slide_data.lead_sentence or ''}"
        p_lead.font.size = Pt(12.5)
        p_lead.font.bold = True
        p_lead.font.color.rgb = COLOR_PRIMARY_DARK

        # コンテンツブロック（カードカラム）のレイアウト計算
        blocks = slide_data.blocks
        num_blocks = len(blocks)
        if num_blocks == 0:
            return

        total_width = 11.73  # Inches
        start_x = 0.8        # Inches
        start_y = 2.3        # Inches
        card_height = 4.3    # Inches
        gap = 0.3            # Inches

        col_width = (total_width - (gap * (num_blocks - 1))) / num_blocks

        for i, block in enumerate(blocks):
            col_x = start_x + i * (col_width + gap)

            # カード背景シェイプ
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(col_x), Inches(start_y), Inches(col_width), Inches(card_height)
            )
            card.fill.solid()
            card.fill.fore_color.rgb = COLOR_CARD_BG
            card.line.color.rgb = COLOR_CARD_BORDER
            card.line.width = Pt(1.0)

            # カードヘッダーのアクセントライン
            header_line = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(col_x + 0.15), Inches(start_y + 0.15), Inches(col_width - 0.3), Inches(0.04)
            )
            header_line.fill.solid()
            header_line.fill.fore_color.rgb = COLOR_ACCENT
            header_line.line.fill.background()

            # カード見出し
            heading_box = slide.shapes.add_textbox(
                Inches(col_x + 0.15), Inches(start_y + 0.25), Inches(col_width - 0.3), Inches(0.55)
            )
            tf_heading = heading_box.text_frame
            tf_heading.word_wrap = True
            tf_heading.margin_left = Inches(0)
            tf_heading.margin_top = Inches(0)
            p_h = tf_heading.paragraphs[0]
            p_h.text = block.heading
            p_h.font.size = Pt(15)
            p_h.font.bold = True
            p_h.font.color.rgb = COLOR_PRIMARY_DARK

            # カード内の箇条書きテキストボックス（見出しと重ならないよう下部に配置）
            body_box = slide.shapes.add_textbox(
                Inches(col_x + 0.15), Inches(start_y + 0.85), Inches(col_width - 0.3), Inches(card_height - 1.0)
            )
            tf_body = body_box.text_frame
            tf_body.word_wrap = True
            tf_body.margin_left = Inches(0)
            tf_body.margin_top = Inches(0)

            # 箇条書き項目数に応じた動的サイズ調整（文字溢れ防止）
            num_points = len(block.bullet_points)
            if num_points >= 4:
                bullet_font_size = Pt(10.5)
                space_after = Pt(6)
                line_spacing = 1.15
            else:
                bullet_font_size = Pt(11.5)
                space_after = Pt(10)
                line_spacing = 1.25

            for j, point in enumerate(block.bullet_points):
                p_b = tf_body.paragraphs[0] if j == 0 else tf_body.add_paragraph()
                p_b.text = f"• {point}"
                p_b.font.size = bullet_font_size
                p_b.font.color.rgb = COLOR_TEXT_BODY
                p_b.space_after = space_after
                p_b.line_spacing = line_spacing

        # スライドフッター注釈
        if slide_data.footer_note:
            ft_box = slide.shapes.add_textbox(
                Inches(0.8), Inches(6.8), Inches(11.73), Inches(0.35)
            )
            tf_ft = ft_box.text_frame
            tf_ft.margin_left = Inches(0)
            tf_ft.margin_top = Inches(0)
            p_ft = tf_ft.paragraphs[0]
            p_ft.text = slide_data.footer_note
            p_ft.font.size = Pt(9.5)
            p_ft.font.color.rgb = COLOR_TEXT_MUTED
