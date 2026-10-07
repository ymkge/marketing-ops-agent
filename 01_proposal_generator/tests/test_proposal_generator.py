"""Unit tests for 01_proposal_generator module."""

import json
import os
import sys
import tempfile
from pathlib import Path
import pytest
from pptx import Presentation

# 01_proposal_generator ディレクトリを sys.path に追加
module_dir = Path(__file__).resolve().parent.parent
if str(module_dir) not in sys.path:
    sys.path.insert(0, str(module_dir))

from src.models import BriefInput, ProposalDeck, SlideContent, ContentBlock
from src.llm_client import generate_mock_deck
from src.pptx_builder import ProposalPptxBuilder


@pytest.fixture
def sample_brief():
    """テスト用企画メモフィクスチャ"""
    return BriefInput(
        client_name="テストクライアント株式会社",
        project_name="新商品プロモーション提案",
        target_service="SaaSサービスA",
        client_challenges=["リード獲得数の伸び悩み", "認知度の不足"],
        product_features=["即日導入可能", "高いROI"],
        target_audience=["マーケティング担当者", "経営企画"],
        budget="月額3,000,000円",
        kpis=["月間リード100件", "CPA 30,000円"],
        strategy=["ホワイトペーパー配信", "リスティング広告"],
        timeline="2026年10月〜12月",
        team_structure="専任ディレクター1名、運用者1名"
    )


def test_brief_model_validation(sample_brief):
    """BriefInputモデルのバリデーションテスト"""
    assert sample_brief.client_name == "テストクライアント株式会社"
    assert len(sample_brief.client_challenges) == 2
    assert sample_brief.budget == "月額3,000,000円"


def test_mock_deck_generation(sample_brief):
    """モックスライド構成生成のテスト（5枚構成の整合性確認）"""
    deck = generate_mock_deck(sample_brief)
    
    assert isinstance(deck, ProposalDeck)
    assert deck.client_name == sample_brief.client_name
    assert len(deck.slides) == 5

    # 5枚のスライド種別が期待通りか
    expected_slide_types = [
        "cover",
        "challenges",
        "target_strategy",
        "simulation",
        "schedule_team"
    ]
    actual_slide_types = [s.slide_type for s in deck.slides]
    assert actual_slide_types == expected_slide_types

    # 各スライドにタイトルとブロックが含まれているか
    for slide in deck.slides:
        assert bool(slide.title)
        assert len(slide.blocks) > 0


def test_pptx_builder_creates_valid_deck(sample_brief):
    """python-pptx によるスライド生成のテスト"""
    deck = generate_mock_deck(sample_brief)
    builder = ProposalPptxBuilder(deck)

    with tempfile.TemporaryDirectory() as tmpdir:
        output_pptx = os.path.join(tmpdir, "test_output.pptx")
        saved_path = builder.build(output_pptx)

        assert os.path.exists(saved_path)
        assert os.path.getsize(saved_path) > 0

        # 生成されたpptxファイルを読み込んで検証
        prs = Presentation(saved_path)
        assert len(prs.slides) == 5
        
        # 16:9比率になっているか確認 (幅 13.333インチ, 高さ 7.5インチ)
        assert round(prs.slide_width.inches, 2) == 13.33
        assert round(prs.slide_height.inches, 2) == 7.50

        # 各スライドにテキストが存在することを確認
        for slide in prs.slides:
            texts = [s.text for s in slide.shapes if s.has_text_frame and s.text]
            assert len(texts) > 0


def test_cli_execution_with_mock():
    """generate_slides.py CLIスクリプトの実行テスト"""
    import subprocess

    sample_json = module_dir / "samples" / "sample_brief.json"
    with tempfile.TemporaryDirectory() as tmpdir:
        output_pptx = os.path.join(tmpdir, "cli_output.pptx")
        cmd = [
            sys.executable,
            str(module_dir / "scripts" / "generate_slides.py"),
            "--input", str(sample_json),
            "--output", output_pptx,
            "--provider", "mock"
        ]
        result = subprocess.run(cmd, capture_output=True, text=True)
        assert result.returncode == 0, f"CLI error: {result.stderr}"
        assert os.path.exists(output_pptx)
        assert os.path.getsize(output_pptx) > 0


def test_cli_execution_with_invalid_input():
    """存在しない入力ファイル指定時のエラーハンドリングテスト"""
    import subprocess

    cmd = [
        sys.executable,
        str(module_dir / "scripts" / "generate_slides.py"),
        "--input", "non_existent_file.json",
        "--provider", "mock"
    ]
    result = subprocess.run(cmd, capture_output=True, text=True)
    assert result.returncode != 0
    assert "見つかりません" in result.stderr or "入力ブリーフの読み込みに失敗しました" in result.stderr


def test_gemini_client_with_mocked_api(sample_brief, monkeypatch):
    """Gemini API をモックして ProposalDeck が生成されることを検証"""
    from unittest.mock import MagicMock
    from src.llm_client import generate_deck_with_gemini

    monkeypatch.setenv("GEMINI_API_KEY", "dummy_gemini_api_key")

    mock_deck = generate_mock_deck(sample_brief)
    mock_response = MagicMock()
    mock_response.text = mock_deck.model_dump_json()

    mock_client = MagicMock()
    mock_client.models.generate_content.return_value = mock_response

    monkeypatch.setattr("google.genai.Client", lambda api_key: mock_client)

    result_deck = generate_deck_with_gemini(sample_brief)
    assert isinstance(result_deck, ProposalDeck)
    assert result_deck.client_name == sample_brief.client_name
    assert len(result_deck.slides) == 5


def test_gemini_missing_api_key(sample_brief, monkeypatch):
    """APIキー未設定時に generate_deck_with_gemini が ValueError を発生させることを検証"""
    from src.llm_client import generate_deck_with_gemini

    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)

    with pytest.raises(ValueError, match="GEMINI_API_KEY または GOOGLE_API_KEY"):
        generate_deck_with_gemini(sample_brief)


def test_unsupported_provider_raises_value_error(sample_brief):
    """未対応のプロバイダー指定時に ValueError が発生することを検証"""
    from src.llm_client import generate_proposal_deck

    with pytest.raises(ValueError, match="未対応のプロバイダー"):
        generate_proposal_deck(sample_brief, provider="invalid_llm")


def test_auto_provider_fallback_to_mock(sample_brief, monkeypatch):
    """APIキーがない場合でも auto モードでモック生成へフォールバックすることを検証"""
    from src.llm_client import generate_proposal_deck

    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    monkeypatch.delenv("GOOGLE_API_KEY", raising=False)
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)

    deck = generate_proposal_deck(sample_brief, provider="auto")
    assert isinstance(deck, ProposalDeck)
    assert len(deck.slides) == 5




